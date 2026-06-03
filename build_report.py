import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from collections import defaultdict
from datetime import datetime

SRC = "Consolidated persona list.xlsx"
OUT = "Consolidated persona list - RESULT.xlsx"
REF_DATE = datetime(2026, 6, 3)   # "today" for days-since calc
STALE_DAYS = 120                  # candidate-for-removal threshold (configurable)

# ---------------------------------------------------------------
# 1) Build lookup maps from a fast read-only pass
# ---------------------------------------------------------------
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

# Tab 3: Persona Name -> set of IDs
ws3 = wb["Personas Key values"]
name_to_ids = defaultdict(list)
for i, row in enumerate(ws3.iter_rows(values_only=True)):
    if i == 0:
        continue
    name, pid = row[0], row[1]
    if name is None or pid is None:
        continue
    nm = str(name).strip()
    kv = str(pid).strip()
    # normalise "724.0" -> "724"
    if kv.endswith(".0"):
        kv = kv[:-2]
    if kv not in name_to_ids[nm]:
        name_to_ids[nm].append(kv)

# Tab 2: GAM -> per key value: total impressions + last date
ws2 = wb["GAM Report"]
def to_int(v):
    if v is None: return 0
    if isinstance(v, (int, float)): return int(v)
    s = str(v).strip().replace(',', '')
    if s in ('', '-'): return 0
    try: return int(float(s))
    except ValueError: return 0

kv_impr = defaultdict(int)
kv_lastdate = {}
for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if i == 0:
        continue
    d, kv, impr = row[0], row[1], row[2]
    if kv is None:
        continue
    kv = str(kv).strip()
    if kv.endswith(".0"):
        kv = kv[:-2]
    kv_impr[kv] += to_int(impr)
    if isinstance(d, datetime):
        if kv not in kv_lastdate or d > kv_lastdate[kv]:
            kv_lastdate[kv] = d

# Tab 1: master persona names (preserve order, keep duplicates)
ws1 = wb["personas_1_2026-05-28.csv"]
tab1 = []
for i, row in enumerate(ws1.iter_rows(values_only=True)):
    if i == 0:
        continue
    name = row[0]
    if name is None or str(name).strip() == "":
        continue
    tab1.append(str(name).strip())
wb.close()

# ---------------------------------------------------------------
# 2) Compute result rows
# ---------------------------------------------------------------
results = []  # dict per persona
for nm in tab1:
    ids = name_to_ids.get(nm, [])
    if not ids:
        results.append(dict(name=nm, ids="", last=None, impr=None, days=None,
                            status="No mapping in Key Values tab",
                            rec="Review - no key value mapped", prio=1))
        continue
    total = 0
    last = None
    any_data = False
    for kv in ids:
        if kv in kv_impr:
            any_data = True
            total += kv_impr[kv]
        d = kv_lastdate.get(kv)
        if d is not None and (last is None or d > last):
            last = d
    if not any_data:
        results.append(dict(name=nm, ids=", ".join(ids), last=None, impr=0, days=None,
                            status="Mapped but no GAM delivery",
                            rec="REMOVE - never delivered (0 impressions)", prio=0))
    else:
        days = (REF_DATE - last).days if last else None
        if last is not None and days > STALE_DAYS:
            status = f"STALE (>{STALE_DAYS}d)"
            rec = f"REMOVE - not used in {days}d"
            prio = 2
        else:
            status = "Active"
            rec = "Keep - active"
            prio = 3
        results.append(dict(name=nm, ids=", ".join(ids), last=last, impr=total,
                            days=days, status=status, rec=rec, prio=prio))

# ---------------------------------------------------------------
# 3) Write output workbook (copy original, fill Tab1, add summary)
# ---------------------------------------------------------------
wb_out = openpyxl.load_workbook(SRC)  # full load to preserve all tabs
ws = wb_out["personas_1_2026-05-28.csv"]

# headers
headers = ["Persona Name", "Last date used", "Overall impressions till date",
           "Key Value(s) (Persona ID)", "Days since last used",
           "Status / Note", "Recommendation"]
hdr_fill = PatternFill("solid", fgColor="1F4E78")
hdr_font = Font(bold=True, color="FFFFFF")
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

stale_fill = PatternFill("solid", fgColor="F8CBAD")     # orange-ish
nomap_fill = PatternFill("solid", fgColor="D9D9D9")     # grey
nodel_fill = PatternFill("solid", fgColor="FFF2CC")     # yellow
NCOL = len(headers)

def style_row(sheet, idx, status):
    if status.startswith("STALE"):
        f = stale_fill
    elif status == "No mapping in Key Values tab":
        f = nomap_fill
    elif status == "Mapped but no GAM delivery":
        f = nodel_fill
    else:
        return
    for c in range(1, NCOL + 1):
        sheet.cell(row=idx, column=c).fill = f

for idx, r in enumerate(results, start=2):
    ws.cell(row=idx, column=1, value=r["name"])
    ws.cell(row=idx, column=2,
            value=r["last"].strftime("%Y-%m-%d") if r["last"] else "")
    ws.cell(row=idx, column=3, value=("" if r["impr"] is None else r["impr"]))
    ws.cell(row=idx, column=4, value=r["ids"])
    ws.cell(row=idx, column=5, value=("" if r["days"] is None else r["days"]))
    ws.cell(row=idx, column=6, value=r["status"])
    ws.cell(row=idx, column=7, value=r["rec"])
    style_row(ws, idx, r["status"])

# clear any leftover rows below our data (original had blank rows up to 1000)
last_written = len(results) + 1
if ws.max_row > last_written:
    ws.delete_rows(last_written + 1, ws.max_row - last_written)

# column widths
widths = [40, 16, 26, 26, 20, 28, 38]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "A2"

# ---------------- Decision Summary sheet (sorted by removal priority) -------
if "Decision Summary" in wb_out.sheetnames:
    del wb_out["Decision Summary"]
summ = wb_out.create_sheet("Decision Summary", 0)

def sort_key(r):
    # removal priority first (0=never delivered ... 3=active),
    # then impressions ascending (lowest volume first)
    return (r["prio"], (r["impr"] if r["impr"] is not None else -1))

ordered = sorted(results, key=sort_key)
for c, h in enumerate(headers, start=1):
    cell = summ.cell(row=1, column=c, value=h)
    cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for idx, r in enumerate(ordered, start=2):
    summ.cell(row=idx, column=1, value=r["name"])
    summ.cell(row=idx, column=2, value=r["last"].strftime("%Y-%m-%d") if r["last"] else "")
    summ.cell(row=idx, column=3, value=("" if r["impr"] is None else r["impr"]))
    summ.cell(row=idx, column=4, value=r["ids"])
    summ.cell(row=idx, column=5, value=("" if r["days"] is None else r["days"]))
    summ.cell(row=idx, column=6, value=r["status"])
    summ.cell(row=idx, column=7, value=r["rec"])
    style_row(summ, idx, r["status"])
for i, w in enumerate(widths, start=1):
    summ.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
summ.freeze_panes = "A2"

# ---------------- Remove List sheet (only removal candidates) ---------------
# Includes: never-delivered (0 impressions), unmapped names, and date-stale rows.
if "Remove List" in wb_out.sheetnames:
    del wb_out["Remove List"]
rem = wb_out.create_sheet("Remove List", 1)

rem_headers = ["Persona Name", "Key Value(s) (Persona ID)", "Last date used",
               "Overall impressions till date", "Reason to remove"]
for c, h in enumerate(rem_headers, start=1):
    cell = rem.cell(row=1, column=c, value=h)
    cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

reason_map = {
    "Mapped but no GAM delivery": "Has a key value but never delivered any impressions",
    "No mapping in Key Values tab": "Not present in the Key Values tab - no key value found",
}
remove_rows = [r for r in ordered if r["prio"] in (0, 1, 2)]
for idx, r in enumerate(remove_rows, start=2):
    reason = reason_map.get(r["status"], f"Stale - last used {r['days']}d ago")
    rem.cell(row=idx, column=1, value=r["name"])
    rem.cell(row=idx, column=2, value=r["ids"])
    rem.cell(row=idx, column=3, value=r["last"].strftime("%Y-%m-%d") if r["last"] else "")
    rem.cell(row=idx, column=4, value=("" if r["impr"] is None else r["impr"]))
    rem.cell(row=idx, column=5, value=reason)
    style_row(rem, idx, r["status"])
rem_widths = [40, 26, 16, 28, 52]
for i, w in enumerate(rem_widths, start=1):
    rem.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
rem.freeze_panes = "A2"

wb_out.save(OUT)
print("Remove List rows:", len(remove_rows))

# ---------------- Standalone Remove List CSV --------------------------------
import csv
with open("Remove List.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(rem_headers)
    for r in remove_rows:
        reason = reason_map.get(r["status"], f"Stale - last used {r['days']}d ago")
        w.writerow([r["name"], r["ids"],
                    r["last"].strftime("%Y-%m-%d") if r["last"] else "",
                    ("" if r["impr"] is None else r["impr"]), reason])
print("Wrote Remove List.csv")

# ---------------------------------------------------------------
# 4) Console summary
# ---------------------------------------------------------------
from collections import Counter
cnt = Counter(r["status"] for r in results)
print("OUTPUT WRITTEN:", OUT)
print("Total persona rows:", len(results))
for k, v in cnt.most_common():
    print(f"  {k}: {v}")
print("\nNever-delivered personas (strong remove candidates):")
nd = [r for r in ordered if r["status"] == "Mapped but no GAM delivery"]
for r in nd:
    print(f"  - {r['name']}  [id {r['ids']}]")
print("\nLowest-volume ACTIVE personas (next removal candidates):")
act = [r for r in ordered if r["status"] == "Active"][:10]
for r in act:
    print(f"  {r['impr']:>14,}  last={r['last'].strftime('%Y-%m-%d')}  {r['name']}  [{r['ids']}]")
